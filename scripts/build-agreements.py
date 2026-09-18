# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx==1.1.2"]
# ///
"""Build the Luwah Technologies client agreement templates.

One content model renders four documents as DOCX, then exports each to PDF
through Microsoft Word over AppleScript on macOS.

    uv run scripts/build-agreements.py                 # writes public/agreements/
    uv run scripts/build-agreements.py --no-pdf        # DOCX only, no Word needed
    uv run scripts/build-agreements.py --copy-to DIR   # also copy outputs to DIR

Brand: Typography_Standards_2025-11-02 (OneDrive, 2026/Brand_Identity). Headings
Tw Cen MT, body Open Sans, US Copper #B87333, Lake Blue #4A90A4.

Vikunja #311 and #312. Not legal advice. Review with counsel before first client use.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
LOGO = ROOT / "public" / "images" / "logo-footer.png"
OUT_DEFAULT = ROOT / "public" / "agreements"

VERSION = "v1.0"
VERSION_DATE = "2026-09-18"
VERSION_LABEL = "September 2026"

COMPANY = "Luwah Technologies LLC"
COMPANY_SHORT = "Luwah Technologies"
STATE = "Colorado"
COUNTY = "Arapahoe County, Colorado"
CITY_LINE = "Aurora, Colorado 80017"
EMAIL = "hello@luwahtechnologies.com"
PHONE = "+1 (720) 421-7184"
WEB = "luwahtechnologies.com"
SIGNER = "Daniel Cooke"
SIGNER_TITLE = "Founder"
TAGLINE = "Where Simplicity Meets Innovation"

COPPER = RGBColor(0xB8, 0x73, 0x33)
BLUE = RGBColor(0x4A, 0x90, 0xA4)
INK = RGBColor(0x26, 0x26, 0x26)
MUTED = RGBColor(0x66, 0x66, 0x66)
SHADE = "F3F6F7"
HEAD_FONT = "Tw Cen MT"
BODY_FONT = "Open Sans"

# --------------------------------------------------------------------------
# Content model
# --------------------------------------------------------------------------


@dataclass
class Section:
    title: str
    blocks: list = field(default_factory=list)


@dataclass
class Doc:
    stem: str
    title: str
    subtitle: str
    summary: list[tuple[str, str]]
    intro: list
    sections: list[Section]
    parties: tuple[str, str]  # column headings for the signature block
    exhibits: list[Section]
    left_signer: list[tuple[str, str]]
    right_signer: list[tuple[str, str]]
    doc_prefix: str = "LT-DOC"  # id family printed on the cover, e.g. [MSA-YYYY-NN]
    cover_tag: str = "CLIENT AGREEMENT"  # family label on the cover, kept to one line
    cover_line: str = "[Client name]   |   [Client website or project]"
    cover_parties: tuple[str, ...] = (f"{COMPANY} (\"Provider\")", "[Client legal name] (\"Client\")")


def C(text: str):
    return ("c", text)


def CL(label: str, text: str):
    return ("cl", label, text)


def CLB(label: str, text: str):
    """Labelled clause whose body is bold. Colorado C.R.S. 8-70-115 wants the
    contractor disclosure larger, bold or underlined, not merely capitalised."""
    return ("clb", label, text)


def B(items: list[str]):
    return ("b", items)


def P(text: str):
    return ("p", text)


def N(text: str):
    return ("n", text)


def H(text: str):
    return ("h", text)


def T(header: list[str], rows: list[list[str]], widths: list[float] | None = None):
    return ("t", header, rows, widths)


def F(fields: list[tuple[str, str]]):
    """Two-column fill-in form table."""
    return ("f", fields)


def SIG(left: str, right: str, left_rows: list[tuple[str, str]], right_rows: list[tuple[str, str]]):
    return ("sig", left, right, left_rows, right_rows)


BLANK_ROWS = [["", "", "", ""], ["", "", "", ""], ["", "", "", ""]]

def luwah_signer(role: str) -> list[tuple[str, str]]:
    return [
        ("Signature", ""),
        ("Name", SIGNER),
        ("Title", SIGNER_TITLE),
        ("Email", EMAIL),
        ("Date", ""),
        ("Notice address", f"[{role} notice address]"),
    ]


PROVIDER_SIGNER = luwah_signer("Provider")

CLIENT_SIGNER = [
    ("Signature", ""),
    ("Name", "[Signer name]"),
    ("Title", "[Signer title]"),
    ("Email", "[Signer email]"),
    ("Date", ""),
    ("Notice address", "[Client notice address]"),
]

AI_TOOLS_CLAUSE = (
    "Provider uses AI-assisted tools, including Anthropic Claude and Google Gemini, where they help, "
    "to build Work Product and to help manage the Services, for example code, documentation, meeting "
    "notes and operations. Provider applies human authorship, review, selection and arrangement to all "
    "Work Product and to every change made while operating it, and remains responsible for the result. "
    "Ownership of Work Product under this Agreement applies regardless of the tools used to create it, "
    "and Provider's warranties are about performance and the right to transfer, not about the method of "
    "authorship. Provider submits Client Confidential Information only to AI tools operated under terms "
    "that do not permit the vendor to train on submitted content or to use it for any purpose other "
    "than providing the service and its safety review, and a submission on those terms is a permitted "
    "disclosure under the confidentiality Section of this Agreement. Provider does not submit "
    "credentials to any AI tool, and submits protected health information, payment card data or other "
    "regulated data only where an SOW or other written agreement between the Parties expressly "
    "provides for it. Where an AI "
    "tool records or transcribes a meeting, Provider tells the participants before it starts. Client "
    "acknowledges that the copyright status of material generated by an AI tool without human "
    "authorship is unsettled, and that Provider does not warrant copyright protection for any such "
    "portion. On request Provider will state which AI tools were used on a Deliverable."
)

# Clauses shared word for word across documents. Hoisted so a change lands everywhere.
DTSA_NOTICE = (
    "Under the Defend Trade Secrets Act, 18 U.S.C. § 1833(b), an individual is not liable under any "
    "federal or state trade secret law for disclosing a trade secret in confidence to a federal, state "
    "or local government official, or to an attorney, solely to report or investigate a suspected "
    "violation of law, or in a complaint or other document filed under seal in a lawsuit. Nothing in "
    "this Agreement limits that immunity or prevents any person from reporting a possible violation of "
    "law to a government agency."
)

GENERAL_KNOWLEDGE_CLAUSE = (
    "Nothing in this Agreement prohibits {party} from using the general training, knowledge, skills and "
    "experience {party} gains through the Services, including methods and techniques of general "
    "application, provided {party} does not use or disclose Confidential Information in doing so. This "
    "clause grants no license under any intellectual property right."
)

NOTICES_CLAUSE = CL("Notices.", "Notices must be in writing and sent by email to the address in the signature block. A notice of breach or termination must also be sent by courier or certified mail to the notice address in the signature block. Email notice is effective on the next business day after sending, unless the sender receives a delivery failure. A Party may change its notice address by notice under this Section.")
SEVERABILITY_CLAUSE = CL("Severability, counterparts and electronic signature.", "If a court finds a term unenforceable, the court will limit it to the minimum extent needed and the rest of this Agreement stays in effect. This Agreement and any document signed under it may be signed in counterparts and electronically. An electronic or scanned signature has the same effect as an original.")
HEADINGS_CLAUSE = CL("Headings and independent advice.", "Headings are for convenience only and do not affect interpretation. Each Party has had the opportunity to obtain independent legal advice before signing.")

ACCESSIBILITY_TOOLKIT_CLAUSE = (
    "The Luwah Accessibility Toolkit, meaning the Accessibility Center launcher and drawer, the "
    "disability profiles and individual adjustments, the page-structure navigator, the reading "
    "overlays, the image description layer, the issue report form, the accessibility statement "
    "content, the keyboard shortcut handling and the CSS hooks, together with the code, styles and "
    "documentation that support them, is Background IP that Provider designed, built and owns. Where "
    "the Toolkit is delivered with a website, Client receives the license in {section} as part of the "
    "Deliverable, on {scope} only. Client may configure and modify the Toolkit on that website. Client "
    "may not copy the Toolkit to another website, application, product or system, remove it for "
    "separate use, sublicense, sell, rent, lease or distribute it on its own, remove or alter "
    "Provider's notices in it, present it as Client's own product or as a third-party vendor's "
    "product, or let a third party do any of those things. Use on any other property needs a separate "
    "written license from Provider. The accessibility statement published with the Toolkit describes "
    "a self-assessment against WCAG 2.2 AA, not a certification, and Client keeps that statement "
    "accurate to the website it describes after delivery."
)

AI_CONFIDENTIALITY_CLAUSE = (
    "A Receiving Party may use an AI-assisted tool, including Anthropic Claude and Google Gemini, "
    "to work with Confidential Information only where the tool is operated under terms that do not "
    "permit the vendor to train on submitted content or to use it for any purpose other than "
    "providing the service and its safety review. A submission on those terms is a permitted "
    "disclosure under this Section. The Receiving Party is responsible for anything such a tool "
    "produces from Confidential Information as if the Receiving Party had produced it. The Receiving "
    "Party does not submit credentials or account passwords to any AI tool, and submits protected "
    "health information, payment card numbers or other regulated data only where a written agreement "
    "between the Parties expressly provides for it. Where an AI tool records or transcribes a meeting, the Party "
    "using it tells the other participants before it starts."
)

TEMPLATE_NOTE = (
    f"Template {VERSION}, {VERSION_DATE}. Replace every bracketed field, fill or delete "
    "every schedule, and delete this note before sending. The signed document governs, "
    "not this template."
)

# --------------------------------------------------------------------------
# Master Services Agreement
# --------------------------------------------------------------------------

MSA = Doc(
    stem="MSA_LuwahTech_Template_v1",
    title="Master Services Agreement",
    subtitle=f"{COMPANY} standard terms for automation, web and technology services",
    summary=[
        ("Agreement No.", "[MSA-YYYY-NN]"),
        ("Client", "[Client legal name and entity type]"),
        ("Client contact", "[Name, title, email]"),
        ("Effective Date", "[Month DD, YYYY]"),
        ("Initial SOW reference", "[SOW-YYYY-NN, or None at signing]"),
        ("Total value", "[Sum of signed SOWs, updated as SOWs are added]"),
        ("Initial Term end", "[Effective Date plus 12 months]"),
    ],
    intro=[
        P(
            f"This Master Services Agreement (the \"Agreement\") is between {COMPANY}, a "
            f"{STATE} limited liability company (\"Provider\"), and the client named in the "
            "Agreement Summary above (\"Client\"). Provider and Client are each a \"Party\" and "
            "together the \"Parties\". The Agreement takes effect on the Effective Date."
        ),
        P(
            "This Agreement sets the terms that apply to every project. Each project is described "
            "in a Statement of Work. Signing a Statement of Work does not reopen these terms."
        ),
    ],
    sections=[
        Section("Definitions", [
            CL("\"Statement of Work\" or \"SOW\"", "means a document in substantially the form of Exhibit A, signed by both Parties, that describes the Services, Deliverables, fees and dates for a project or a recurring service."),
            CL("\"Services\"", "means the work Provider performs under an SOW."),
            CL("\"Deliverables\"", "means the items an SOW identifies for delivery to Client."),
            CL("\"Work Product\"", "means the Deliverables and any other original material Provider creates specifically for Client under an SOW. Work Product excludes Background IP."),
            CL("\"Background IP\"", "means tools, code, libraries, templates, workflows, methods, know-how and documentation that Provider created or acquired outside the Services, or creates during the Services for general use rather than specifically for Client, and any improvements to them. Exhibit B lists the Background IP known at signing. Background IP that is not listed is still Background IP."),
            CL("\"Client Materials\"", "means the data, content, credentials, accounts, systems and other material Client provides to Provider."),
            CL("\"Confidential Information\"", "means non-public information one Party discloses to the other, in any form, that is marked confidential or that a reasonable person would treat as confidential. It includes business data, credentials, pricing, customer lists, source code and the terms of this Agreement."),
            CL("\"Change Order\"", "means a document in substantially the form of Exhibit C, signed by both Parties."),
            CL("\"Acceptance\" and \"Project Completion\"", "have the meanings in Section 10."),
        ]),
        Section("How This Agreement and SOWs Work Together", [
            C("This Agreement applies to every SOW. Each SOW incorporates this Agreement. Together they form one contract for that SOW."),
            C("An SOW may change a term of this Agreement only by naming the section it changes. The change applies to that SOW only."),
            C("If the Parties signed a proposal, letter of agreement or quote for the same work, this Agreement and the SOW control over it."),
            C("Each SOW is a separate obligation. Ending one SOW does not end this Agreement or any other SOW."),
        ]),
        Section("Services", [
            C("Provider will perform the Services and deliver the Deliverables described in each SOW with reasonable skill and care, consistent with industry standards."),
            C("Where an SOW defines a Deliverable by its outcome, Provider chooses the method, tools and platforms used to reach that outcome. A change to a third-party platform, API or service that requires a different method is not a change in scope."),
            C("Work not described in an SOW is out of scope. Out-of-scope work needs a signed Change Order before it starts. Provider may decline a change request."),
            C("Provider may use subcontractors. Provider remains responsible for their work and binds each of them in writing to confidentiality and intellectual property terms at least as protective as this Agreement."),
            C(AI_TOOLS_CLAUSE),
        ]),
        Section("Client Responsibilities", [
            C("Client will provide, at its own cost, the access, information, decisions and approvals each SOW requires, and will respond to Provider's requests within 3 business days unless the SOW says otherwise."),
            C("Client will name one point of contact with authority to approve Deliverables and Change Orders."),
            C("Client is responsible for the accuracy and lawful use of Client Materials, for holding the rights needed to give them to Provider, and for backups of Client data before, during and after the Services, unless an SOW assigns backups to Provider."),
            C("Delay caused by Client extends every affected date by at least the length of the delay. If Client delay exceeds 30 days, Provider may invoice the work completed to date and pause the SOW until Client is ready to continue."),
        ]),
        Section("Fees and Payment", [
            C("Client pays the fees stated in each SOW. All amounts are in US dollars."),
            C("Unless the SOW says otherwise, a deposit of 50% of any fixed fee is due when the SOW is signed. Work starts when the SOW is signed and the deposit has cleared."),
            C("Hourly work is billed in 15-minute increments at the rate schedule in the SOW. Provider invoices hourly work monthly with an itemized summary, unless the SOW says otherwise."),
            C("Invoices are due within 14 days of the invoice date (Net 14). Client pays by ACH, credit card, debit card, or another method Provider accepts in writing."),
            C("Overdue amounts accrue a late charge of 1.5% per month, or the highest rate the law allows if that is lower. If an invoice is more than 10 days overdue, Provider may suspend the Services and withhold Deliverables after 5 days' written notice, until payment clears."),
            C("Client must dispute an invoice in writing within 7 days of the invoice date and state the reason. Client pays the undisputed part on time. The Parties resolve the dispute in good faith."),
            C("Third-party costs an SOW requires, such as hosting, software licenses, domain registrations and API usage, are pass-through costs. Client pays them directly where possible. Where Provider pays them, Provider invoices them at cost with receipts, and obtains Client's prior written approval for any single item over $100."),
            C("Fees exclude taxes. Client pays any sales, use or similar tax that applies to the Services, other than tax on Provider's income."),
            C("Provider applies each deposit to the amounts due under the SOW. Deposits are not refunded once work has started, except as Section 6.3 provides."),
        ]),
        Section("Term and Termination", [
            C("This Agreement starts on the Effective Date and continues for 12 months (the \"Initial Term\"). It then renews for successive 12-month terms unless either Party gives written notice of non-renewal at least 30 days before the current term ends. Expiry of this Agreement does not end an SOW already in progress. This Agreement continues to govern that SOW until it is complete or terminated."),
            C("An SOW for a recurring service, such as hosting, maintenance or a retainer, runs month to month after any initial period stated in the SOW. Either Party may end it with 30 days' written notice."),
            C("Either Party may terminate an SOW for convenience with 14 days' written notice. Client pays for all Services performed and expenses incurred up to the termination date, including work in progress, at the SOW rates. For a fixed fee, Client pays the greater of the deposit already paid and a proportionate share of the fee for the work completed. Provider refunds any part of the deposit above that amount, and Client pays any shortfall within 14 days."),
            C("Either Party may terminate this Agreement or an SOW for material breach if the breach is not cured within 10 days of written notice describing it. Non-payment is a material breach."),
            C("When an SOW ends for any reason:"),
            B([
                "Client pays all amounts due within 14 days.",
                "On receipt of payment, Provider delivers the Work Product completed to that date.",
                "Provider returns or deletes Client Materials and Client Confidential Information within 30 days, except copies held in routine backups or that the law requires Provider to keep.",
                "Client revokes Provider's access to Client systems, and Provider deletes any Client credentials it holds, within 5 business days.",
                "If Client asks, Provider gives reasonable transition assistance at the hourly rate in the SOW or Provider's then-current rate.",
            ]),
            C("Sections 5, 7, 8, 9, 10.3 to 10.5, 11, 12, 14 and 15, and any other term that by its nature should survive, survive the end of this Agreement."),
        ]),
        Section("Intellectual Property", [
            CL("Client Materials.", "Client keeps all rights in Client Materials. Client grants Provider a non-exclusive license to use Client Materials only to perform the Services."),
            CL("Work Product.", "When Client has paid in full all fees due under the applicable SOW, the Work Product is a \"work made for hire\" for Client to the extent the law allows. To the extent any Work Product is not a work made for hire, Provider assigns to Client, on that payment, all right, title and interest in it worldwide, including copyright. Until payment in full, Provider owns the Work Product and Client may use it only to review it."),
            CL("Background IP.", "Provider keeps all rights in Background IP. Where Background IP is incorporated into a Deliverable, Provider grants Client a perpetual, worldwide, non-exclusive, royalty-free license to use, copy, modify and distribute that Background IP as part of the Deliverable for Client's business purposes, except as Section 7.7 provides. Client may not sell or license the Background IP on its own or separate it from the Deliverable for that purpose."),
            CL("Third-party and open source components.", "Deliverables may include third-party or open source components. Those components remain subject to their own licenses, which Provider will identify on request. Provider will not include a component whose license would require Client to publish Client's own source code without Client's written approval."),
            CL("General knowledge.", "Nothing in this Agreement stops Provider from using the general skills, knowledge, experience and techniques it gains while performing the Services, provided Provider does not use or disclose Client Confidential Information."),
            CL("Portfolio.", "Provider may name Client and describe the general nature of the Services in Provider's portfolio and marketing only with Client's prior written consent, which may be given by email, and without disclosing Confidential Information. Client may withdraw consent at any time by written notice."),
            CL("Accessibility toolkit.", ACCESSIBILITY_TOOLKIT_CLAUSE.format(section="Section 7.3", scope="the websites and systems the SOW names")),
            CL("Further assurances.", "Each Party will sign the documents reasonably needed to record the rights granted in this Section."),
        ]),
        Section("Confidentiality", [
            C("Each Party will use the other Party's Confidential Information only to perform or receive the Services. Each Party will protect it with at least the care it uses for its own confidential information, and never less than reasonable care. Each Party will disclose it only to its employees, contractors and advisers who need to know it and who are bound by written confidentiality obligations at least as protective as this Section."),
            C("Confidential Information does not include information that (a) is or becomes public through no fault of the receiving Party, (b) the receiving Party already knew without a duty of confidence, (c) the receiving Party lawfully receives from a third party without a duty of confidence, or (d) the receiving Party develops independently without using the other Party's Confidential Information."),
            C("A Party may disclose Confidential Information when the law or a court requires it, after giving the other Party prompt notice where lawful and cooperating with any effort to limit the disclosure."),
            C("These obligations last for 3 years after this Agreement ends. For trade secrets and credentials, they last as long as the information remains a trade secret or the credential remains active."),
            C("If a mutual non-disclosure agreement between the Parties is in force, this Section applies in addition to it. Where they conflict, the term that gives the disclosing Party more protection applies."),
        ]),
        Section("Credentials, Data and Security", [
            C("Client will grant Provider the least access needed for the Services, through named accounts or roles where the platform allows, rather than shared passwords."),
            C("Provider stores Client credentials in an encrypted password manager, enables multi-factor authentication where the platform offers it, does not share credentials with anyone not working on the Services, and does not store them in email, chat or source code."),
            C("Provider will notify Client within 72 hours of learning of unauthorized access to Client systems or data that arises from the Services, and will cooperate with Client's response."),
            C("Client remains responsible for its own data and for the legal obligations that apply to it, including any privacy or sector-specific law. Provider will follow Client's reasonable written instructions about Client data. Where the Services involve regulated data, the SOW must say so and may add terms."),
            C("Provider is not responsible for the availability, performance or changes of third-party platforms, services or APIs, or for problems caused by changes Client or a third party makes to a Deliverable after Acceptance."),
            C("Access ends with the SOW. Section 6.5 applies."),
        ]),
        Section("Acceptance and Warranty", [
            CL("Acceptance.", "When Provider delivers a Deliverable, Client has 7 business days (the \"Review Period\") to test it against the acceptance criteria in the SOW and to report any material non-conformance in writing. Provider corrects each reported non-conformance and redelivers, and a new Review Period of 5 business days starts. A Deliverable is accepted (\"Acceptance\") on the earlier of Client's written acceptance or the end of a Review Period with no written report. Client's use of a Deliverable for its business before the Review Period ends is also Acceptance, except use needed to test it and operation that is part of delivery, such as a Deliverable Provider deploys to Client's live systems in order to deliver it."),
            CL("Project Completion.", "Acceptance of the final Deliverable in an SOW completes that SOW (\"Project Completion\"). The final invoice for fixed-fee work is issued on Project Completion unless the SOW says otherwise."),
            CL("Warranty.", "For 30 days after Acceptance, Provider warrants that each Deliverable will materially conform to the SOW. Client's sole remedy is that Provider will correct the non-conformance at no charge within a reasonable time or, if Provider cannot, refund the fee paid for that Deliverable. This warranty does not cover problems caused by Client Materials, by changes Client or a third party makes, by third-party platform changes, or by use outside the SOW."),
            CL("Authority and non-infringement.", "Provider warrants that it has the right to enter this Agreement and that, to its knowledge, the Work Product will not infringe a third party's intellectual property rights."),
            CL("Disclaimer.", "Except as stated in this Section, Provider provides the Services and Deliverables as is and disclaims all other warranties, express or implied, including merchantability, fitness for a particular purpose and non-infringement. Provider does not warrant that a Deliverable will run without interruption or error, or that it will produce any business result, revenue or saving. Recommendations are advisory. Client decides whether to act on them."),
        ]),
        Section("Limitation of Liability", [
            C("Neither Party is liable to the other for indirect, incidental, special, consequential or punitive damages, or for lost profits, revenue, data or business opportunity, however caused, even if told they were possible."),
            C("Each Party's total liability arising out of or related to an SOW is limited to the fees Client paid to Provider under that SOW in the 12 months before the event that gave rise to the claim."),
            C("The limits in this Section do not apply to (a) a Party's breach of Section 8, (b) Client's payment obligations, (c) a Party's gross negligence, fraud or willful misconduct, (d) a Party's infringement of the other Party's intellectual property, or (e) a Party's obligations under Section 12."),
            C("Client is responsible for backups under Section 4.3. Provider is not liable for data loss that a current backup would have prevented."),
        ]),
        Section("Indemnity", [
            CL("By Provider.", "Provider will defend Client against any third-party claim that Work Product, as delivered by Provider and used as the SOW permits, infringes that third party's copyright, trademark or trade secret, and will pay the damages and costs a court finally awards or a settlement Provider agrees to. This does not cover claims arising from Client Materials, from third-party or open source components used as their licenses allow, from changes Client or a third party makes, or from combining the Work Product with something Provider did not supply. If a claim arises, Provider may modify or replace the Work Product so it no longer infringes, obtain a license, or refund the fee paid for the affected Deliverable and end the license to it."),
            CL("By Client.", "Client will defend Provider against any third-party claim arising from Client Materials, from Client's use of the Work Product in a way the SOW does not permit, or from Client's breach of law, and will pay the damages and costs a court finally awards or a settlement Client agrees to."),
            CL("Process.", "The Party seeking defence must give the other prompt written notice of the claim, give it control of the defence and settlement, and cooperate at the defending Party's expense. The defending Party may not settle in a way that admits fault on the other Party's behalf or imposes an obligation on it without its written consent."),
        ]),
        Section("Relationship and Insurance", [
            C("Provider is an independent contractor. Nothing in this Agreement creates a partnership, joint venture, employment or agency relationship. Neither Party may bind the other."),
            C("Provider is not exclusive to Client and may provide services to others, including Client's competitors, subject to Section 8."),
            C("Provider maintains professional liability (errors and omissions) and cyber liability insurance during the term and will provide a certificate of insurance on request."),
        ]),
        Section("Governing Law and Disputes", [
            C(f"The law of the State of {STATE} governs this Agreement, without regard to its conflict of law rules."),
            C("The Parties will first try to resolve any dispute through good-faith negotiation between their principals for 30 days after one Party gives written notice of the dispute."),
            C(f"If negotiation fails, the Parties will mediate in {COUNTY}, or by video conference, with a mediator they agree on, and will share the mediator's fees equally."),
            C(f"If mediation fails, either Party may bring the dispute in the state or federal courts located in {COUNTY}, and each Party submits to their jurisdiction. Either Party may seek an injunction in any court to protect its Confidential Information or intellectual property without first negotiating or mediating."),
            C("The prevailing Party in any action to enforce this Agreement may recover its reasonable attorneys' fees and costs."),
        ]),
        Section("General", [
            NOTICES_CLAUSE,
            CL("Entire agreement.", "This Agreement, its Exhibits and each SOW are the entire agreement between the Parties about their subject matter. They replace all earlier proposals, discussions and agreements about it, except a mutual non-disclosure agreement, a Work Made for Hire Agreement or another non-disclosure agreement the Parties signed, each of which continues to govern as it provides."),
            CL("Amendment and waiver.", "Changes must be in writing and signed by both Parties. A signature by email or through an e-signature platform counts. A Party's failure to enforce a term is not a waiver of it."),
            CL("Assignment.", "Neither Party may assign this Agreement without the other Party's written consent, except to a successor to substantially all of its business, on written notice."),
            CL("Force majeure.", "Neither Party is liable for delay or failure caused by events beyond its reasonable control, including natural disaster, epidemic, government action, war, civil unrest, labor dispute, or failure of the internet, a utility or a third-party platform. The affected Party will notify the other promptly and resume performance as soon as it can. Payment for work already performed is not excused."),
            SEVERABILITY_CLAUSE,
            HEADINGS_CLAUSE,
        ]),
    ],
    parties=("Provider", "Client"),
    left_signer=PROVIDER_SIGNER,
    right_signer=CLIENT_SIGNER,
    doc_prefix="MSA",
    cover_tag="CLIENT AGREEMENT   ·   MASTER SERVICES",
    exhibits=[
        Section("Exhibit A: Statement of Work", [
            N("Use one SOW per project or recurring service. Every SOW incorporates the Master Services Agreement. Keep the fields below in this order so the contract tracker can read them."),
            F([
                ("SOW No.", "[SOW-YYYY-NN]"),
                ("Client", "[Client legal name]"),
                ("MSA reference", "[MSA-YYYY-NN dated Month DD, YYYY]"),
                ("SOW Effective Date", "[Month DD, YYYY]"),
                ("Start Date", "[Month DD, YYYY]"),
                ("Target Completion Date", "[Month DD, YYYY]"),
                ("End Date or Term", "[Fixed end date, or month to month after an initial period, with the notice period]"),
                ("Total value", "[Fixed fees plus estimated hourly fees, excluding pass-through costs]"),
            ]),
            H("1. Objective"),
            P("[One paragraph. What Client wants to be true when this SOW is complete, and why.]"),
            H("2. Scope of Services"),
            B(["[In-scope item]", "[In-scope item]", "[In-scope item]"]),
            H("3. Out of Scope"),
            B(["[Item that is excluded and would need a Change Order]", "[Item that is excluded]"]),
            H("4. Deliverables and Acceptance Criteria"),
            N("Write each acceptance criterion as a test a reasonable person can run and pass or fail. Acceptance of the last Deliverable is Project Completion under Section 10.2 of the Agreement."),
            T(["No.", "Deliverable", "Acceptance criteria", "Target date"], BLANK_ROWS, [0.5, 2.2, 2.8, 1.0]),
            H("5. Fees and Payment Schedule"),
            T(["Milestone", "Trigger", "Amount (USD)", "Due"], [
                ["Deposit", "Signature of this SOW", "[50% of fixed fee]", "On signing"],
                ["[Milestone]", "[Acceptance of Deliverable N]", "[Amount]", "[Net 14 from invoice]"],
                ["Final", "Project Completion", "[Balance]", "Net 14 from invoice"],
            ], [1.5, 2.3, 1.4, 1.3]),
            H("6. Rate Schedule"),
            N("Applies to hourly work, support after the warranty period, and Change Orders unless a Change Order states a fixed price."),
            T(["Service", "Rate (USD)", "Unit"], [
                ["[Development and automation]", "[Rate]", "Hour, 15-minute increments"],
                ["[Consulting and training]", "[Rate]", "Hour, 15-minute increments"],
                ["[Support and maintenance retainer]", "[Monthly fee]", "Month, includes [N] hours"],
            ], [3.0, 1.5, 2.0]),
            H("7. Third-Party Costs"),
            T(["Item", "Paid by", "Estimated cost (USD)", "Notes"], [
                ["[Hosting]", "[Client directly]", "[Amount per month]", ""],
                ["[Domain, licenses, API usage]", "[Client directly, or Provider at cost]", "[Amount]", ""],
            ], [2.0, 1.8, 1.5, 1.2]),
            H("8. Client Responsibilities and Dependencies"),
            B(["[Access, accounts or credentials Client must provide, and by when]", "[Content, data or decisions Client must provide, and by when]", "[Third-party approvals the project depends on]"]),
            H("9. Assumptions"),
            B(["[Assumption the estimate relies on]", "[Assumption the estimate relies on]"]),
            H("10. Terms That Differ From the Agreement"),
            N("Name the section of the Agreement and state the change. Leave blank if there is none."),
            T(["Agreement section", "Change for this SOW"], [["", ""], ["", ""]], [2.0, 4.5]),
            SIG("Provider", "Client", PROVIDER_SIGNER[:5], CLIENT_SIGNER[:5]),
        ]),
        Section("Exhibit B: Background IP Schedule", [
            N("List every item of Provider Background IP expected to be incorporated in Deliverables. Add items for a later SOW by Change Order. This schedule is mandatory. If nothing applies, write NONE in the first row. The example rows are Provider's usual reusable assets and must be edited before sending."),
            T(["Item", "Description", "Where it is used", "License to Client"], [
                ["n8n workflow library", "Reusable sub-workflows and patterns for error handling, notification, deduplication and scheduling", "Automation Deliverables", "Section 7.3"],
                ["Web starter and components", "Next.js site starter, layout components, design tokens, forms and deployment configuration", "Website Deliverables", "Section 7.3"],
                ["Luwah accessibility toolkit", "Accessibility launcher and panel, disability profiles, adjustments, reading overlays, page-structure navigator, issue report form, statement seed content and the supporting code and styles. Licensed to the delivered website only, Section 7.7", "Website Deliverables", "Section 7.3"],
                ["Operations scripts", "Backup, monitoring, deployment and health-check scripts and templates", "Hosting and maintenance Deliverables", "Section 7.3"],
                ["Documentation templates", "Runbook, handover and training document templates", "All Deliverables", "Section 7.3"],
                ["[Item]", "[Description]", "[Where used]", "Section 7.3"],
            ], [1.6, 2.6, 1.4, 0.9]),
        ]),
        Section("Exhibit C: Change Order", [
            N("Use a Change Order for any work outside a signed SOW, or to change a Deliverable, date or fee. Work on the change starts only after both Parties sign."),
            F([
                ("Change Order No.", "[CO-YYYY-NN]"),
                ("SOW No.", "[SOW-YYYY-NN]"),
                ("Client", "[Client legal name]"),
                ("Date", "[Month DD, YYYY]"),
                ("Requested by", "[Name and Party]"),
            ]),
            H("1. Description of the Change"),
            P("[What changes, and why.]"),
            H("2. Effect on Deliverables and Scope"),
            T(["Deliverable", "Change", "New acceptance criteria"], [["", "", ""], ["", "", ""]], [1.8, 2.4, 2.3]),
            H("3. Effect on Fees"),
            T(["Item", "Basis", "Amount (USD)", "Due"], [["", "[Fixed, or hours at the SOW rate]", "", "[Net 14 from invoice]"], ["Revised total value of SOW", "", "[New total]", ""]], [2.0, 2.0, 1.3, 1.2]),
            H("4. Effect on Schedule"),
            T(["Date", "Was", "Becomes"], [["Target Completion Date", "", ""], ["[Other milestone]", "", ""]], [2.2, 2.1, 2.2]),
            H("5. Effect on Background IP"),
            N("Add any Provider Background IP this change introduces, in the same form as Exhibit B. Write NONE if nothing changes."),
            T(["Item", "Description", "Where it is used", "License to Client"], [["[Item, or NONE]", "", "", "Section 7.3"]], [1.6, 2.6, 1.4, 0.9]),
            SIG("Provider", "Client", PROVIDER_SIGNER[:5], CLIENT_SIGNER[:5]),
        ]),
    ],
)

# --------------------------------------------------------------------------
# Work Made for Hire Agreement (client-facing)
# --------------------------------------------------------------------------

WFH = Doc(
    stem="Work_For_Hire_LuwahTech_Template_v1",
    title="Work Made for Hire Agreement",
    subtitle=f"Transfer of ownership in work {COMPANY_SHORT} creates for a client",
    summary=[
        ("Agreement No.", "[WFH-YYYY-NN]"),
        ("Client", "[Client legal name and entity type]"),
        ("Project", "[Project name]"),
        ("Related documents", "[MSA-YYYY-NN and SOW-YYYY-NN, or proposal dated Month DD, YYYY, or None]"),
        ("Effective Date", "[Month DD, YYYY]"),
        ("Total value", "[Total fee for the Project]"),
    ],
    intro=[
        P(
            "This Work Made for Hire Agreement (the \"Agreement\") is between the client named in "
            f"the Agreement Summary above (\"Client\") and {COMPANY}, a {STATE} limited liability "
            "company (\"Provider\"). Client has engaged Provider to create the work described in "
            "Exhibit A (the \"Project\"). The Parties intend that Client will own the results of the "
            "Project on the terms below."
        ),
        P(
            "Use this Agreement when a client wants a standalone, project-specific transfer of "
            "ownership, or when no Master Services Agreement is in place. Where a Master Services "
            "Agreement is in place, this Agreement controls ownership of the Project's Work Product "
            "and the Master Services Agreement controls everything else."
        ),
    ],
    sections=[
        Section("Definitions", [
            CL("\"Work Product\"", "means all deliverables, code, workflows, configurations, designs, content, documentation and other material Provider creates for Client in the Project, including drafts and versions, but excluding Background IP, Client Materials and third-party components."),
            CL("\"Background IP\"", "means tools, code, libraries, templates, workflows, methods, know-how and documentation that Provider created or acquired outside the Project, or creates during the Project for general use rather than specifically for Client, and any improvements to them. Exhibit B lists the Background IP known at signing. Background IP that is not listed is still Background IP."),
            CL("\"Client Materials\"", "means the data, content, credentials, accounts, systems and other material Client provides to Provider."),
            CL("\"Confidential Information\"", "means non-public information one Party discloses to the other that is marked confidential or that a reasonable person would treat as confidential."),
        ]),
        Section("Project and Compensation", [
            C("Provider will create the Work Product described in Exhibit A."),
            C("Client pays the fee stated in Exhibit A on the schedule stated there. If a Master Services Agreement and Statement of Work cover the Project, their payment terms apply instead."),
            C("Ownership transfers under Section 3 only when Client has paid in full."),
        ]),
        Section("Ownership of Work Product", [
            CL("Work made for hire.", "The Parties agree that the Work Product is specially ordered or commissioned by Client and is a \"work made for hire\" under the United States Copyright Act, 17 U.S.C. § 101, to the fullest extent the law allows."),
            CL("Assignment.", "To the extent any Work Product does not qualify as a work made for hire, Provider irrevocably assigns to Client, effective when Client pays in full, all right, title and interest in the Work Product worldwide, including all copyrights, and any patent, trademark and trade secret rights in it, together with the right to sue for past and future infringement."),
            CL("Moral rights.", "To the extent the law allows, Provider waives, and agrees not to assert against Client or its successors, any moral rights in the Work Product."),
            CL("Further assurances.", "Provider will sign any document and take any lawful step Client reasonably requests, at Client's expense, to record or perfect Client's ownership."),
            CL("Before payment in full.", "Until Client pays in full, Provider owns the Work Product and Client may use it only to review it. If Client does not pay in full, no transfer occurs, and Client must stop using and delete the Work Product on Provider's written request."),
            CL("Client's use after transfer.", "After the transfer, Client may use, modify, and distribute the Work Product without restriction, subject to Section 4 for Background IP and third-party components."),
        ]),
        Section("Background IP and Third-Party Materials", [
            CL("Reservation.", "Provider keeps ownership of all Background IP, including the items listed in Exhibit B."),
            CL("License.", "Where Background IP is incorporated into the Work Product, Provider grants Client a perpetual, worldwide, non-exclusive, royalty-free license, irrevocable once Client has paid in full, to use, copy, modify and distribute that Background IP as part of the Work Product for Client's business purposes, except as Section 4.5 provides. Client may transfer this license with the Work Product to a successor or acquirer of Client's business. Client may not sell or license the Background IP on its own."),
            CL("Open source and third-party components.", "The Work Product may include open source or third-party components. Those components remain subject to their own licenses, which Provider will identify on request. Provider will not include a component whose license would require Client to publish Client's own source code without Client's written approval."),
            CL("Reuse.", "Provider may reuse its Background IP, and the general skills, knowledge, patterns and techniques it gains in the Project, for other clients, provided Provider does not use or disclose Client Confidential Information or the Work Product itself."),
            CL("Accessibility toolkit.", ACCESSIBILITY_TOOLKIT_CLAUSE.format(section="Section 4.2", scope="the Project's website")),
        ]),
        Section("Client Materials", [
            C("Client keeps all rights in Client Materials. Client grants Provider a non-exclusive license to use Client Materials only to perform the Project."),
            C("Client warrants that it has the rights needed to give the Client Materials to Provider for the Project and that their use in the Project will not infringe a third party's rights."),
        ]),
        Section("Provider's Representations", [
            C("Provider has the right to enter this Agreement and to grant the rights in it, and no other agreement conflicts with it."),
            C("The Work Product is Provider's original work, except for Background IP, Client Materials and third-party components identified under Section 4.3, and to Provider's knowledge does not infringe any third party's intellectual property rights."),
            C("Provider has obtained, or will obtain before delivery, from any subcontractor who contributes to the Work Product a written assignment sufficient to give Client the rights in this Agreement."),
            C(AI_TOOLS_CLAUSE),
        ]),
        Section("Confidentiality", [
            C("Each Party will use the other Party's Confidential Information only for the Project, will protect it with reasonable care, and will not disclose it to anyone other than its employees, contractors and advisers who need to know it and are bound by written confidentiality obligations."),
            C("These obligations do not apply to information that is or becomes public through no fault of the receiving Party, that the receiving Party already knew or lawfully receives without a duty of confidence, or that the receiving Party develops independently. A Party may disclose Confidential Information when the law requires, after prompt notice where lawful."),
            C("These obligations last for 3 years after the Project ends, and for trade secrets and credentials for as long as they remain a trade secret or active. If a non-disclosure agreement or Master Services Agreement between the Parties is in force, its confidentiality terms apply in addition, and the more protective term controls."),
        ]),
        Section("Portfolio and Credit", [
            C("Provider may name Client and describe the general nature of the Project in Provider's portfolio and marketing only with Client's prior written consent, which may be given by email, and without disclosing Confidential Information. Client may withdraw consent at any time by written notice."),
        ]),
        Section("Relationship of the Parties", [
            C("Provider is an independent contractor. Provider controls the manner and means of performing the Project. Nothing in this Agreement creates a partnership, joint venture, employment or agency relationship, and neither Party may bind the other."),
        ]),
        Section("Warranty and Limitation of Liability", [
            C("If a Master Services Agreement between the Parties covers the Project, its warranty and limitation of liability terms apply and this Section does not."),
            C("Otherwise, for 30 days after delivery, Provider warrants that the Work Product will materially conform to Exhibit A. Client's sole remedy is that Provider will correct the non-conformance at no charge within a reasonable time or, if Provider cannot, refund the fee paid for the non-conforming part. Except as stated in this Section and Section 6, Provider provides the Work Product as is and disclaims all other warranties, express or implied, including merchantability, fitness for a particular purpose and non-infringement."),
            C("Otherwise, Provider will defend Client against any third-party claim that the Work Product, as delivered and used as this Agreement permits, infringes that third party's copyright, trademark or trade secret, and Client will defend Provider against any third-party claim arising from Client Materials or from Client's use of the Work Product in a way this Agreement does not permit. In each case the defending Party pays the damages and costs a court finally awards or a settlement it agrees to, the other Party gives prompt notice and control of the defence, and Provider's defence does not cover Client Materials, third-party components used as their licenses allow, or changes made by anyone other than Provider."),
            C("Neither Party is liable to the other for indirect, incidental, special, consequential or punitive damages, or for lost profits, revenue, data or business opportunity. Each Party's total liability arising out of this Agreement is limited to the fees paid under it. These limits do not apply to a breach of Section 7, to Client's payment obligations, to a Party's defence obligations under Section 10.3, or to a Party's gross negligence, fraud or willful misconduct."),
        ]),
        Section("Term and Termination", [
            C("This Agreement starts on the Effective Date and ends when the Project is complete and Client has paid in full, unless ended earlier under this Section."),
            C("Either Party may terminate this Agreement with 14 days' written notice. Client pays for all work performed to the termination date. Section 3 applies to the Work Product completed to that date once that payment is made."),
            C("Sections 3, 4, 6, 7, 8, 10 and 12 survive termination."),
        ]),
        Section("General", [
            CL("Governing law and disputes.", f"The law of the State of {STATE} governs this Agreement. The Parties will first negotiate any dispute in good faith for 30 days, then mediate in {COUNTY} or by video conference with a mediator they agree on, sharing the mediator's fees equally. If mediation fails, either Party may bring the dispute in the state or federal courts located in {COUNTY}, and each Party submits to their jurisdiction. Either Party may seek an injunction in any court to protect its Confidential Information or intellectual property without first negotiating or mediating. The prevailing Party may recover its reasonable attorneys' fees and costs."),
            NOTICES_CLAUSE,
            CL("Entire agreement and precedence.", "This Agreement and its Exhibits are the entire agreement between the Parties about ownership of the Work Product. Where a Master Services Agreement or Statement of Work between the Parties also covers the Project, this Agreement controls ownership of the Project's Work Product, and those documents control all other matters."),
            CL("Amendment and assignment.", "Changes must be in writing and signed by both Parties. Provider may not assign this Agreement without Client's written consent. Client may assign this Agreement, with the Work Product, to a successor or acquirer of its business on written notice."),
            SEVERABILITY_CLAUSE,
            HEADINGS_CLAUSE,
        ]),
    ],
    parties=("Provider", "Client"),
    left_signer=PROVIDER_SIGNER,
    right_signer=CLIENT_SIGNER,
    doc_prefix="WFH",
    cover_tag="CLIENT AGREEMENT   ·   WORK MADE FOR HIRE",
    exhibits=[
        Section("Exhibit A: Project Description", [
            N("Describe the Project in enough detail that a third party could tell what is and is not Work Product. Where a Statement of Work exists, reference it and attach it rather than repeating it."),
            F([
                ("Project name", "[Project name]"),
                ("Related SOW or proposal", "[SOW-YYYY-NN, or proposal dated Month DD, YYYY, or None]"),
                ("Description", "[What Provider will create, for what purpose, on which platforms]"),
                ("Delivery date", "[Month DD, YYYY]"),
                ("Total fee (USD)", "[Amount]"),
                ("Payment schedule", "[50% on signing, 50% on delivery, or per the SOW]"),
            ]),
            H("Deliverables"),
            T(["No.", "Deliverable", "Description", "Delivery date"], BLANK_ROWS, [0.5, 2.0, 3.0, 1.0]),
        ]),
        Section("Exhibit B: Provider Background IP", [
            N("List every item of Provider Background IP expected to be incorporated in the Work Product. This schedule is mandatory. If nothing applies, write NONE in the first row. The example rows are Provider's usual reusable assets and must be edited before sending."),
            T(["Item", "Description", "Where it is used", "License to Client"], [
                ["n8n workflow library", "Reusable sub-workflows and patterns for error handling, notification, deduplication and scheduling", "Automation deliverables", "Section 4.2"],
                ["Web starter and components", "Next.js site starter, layout components, design tokens, forms and deployment configuration", "Website deliverables", "Section 4.2"],
                ["Luwah accessibility toolkit", "Accessibility launcher and panel, disability profiles, adjustments, reading overlays, page-structure navigator, issue report form, statement seed content and the supporting code and styles. Licensed to the delivered website only, Section 4.5", "Website deliverables", "Section 4.2"],
                ["Operations scripts", "Backup, monitoring, deployment and health-check scripts and templates", "Hosting and maintenance deliverables", "Section 4.2"],
                ["[Item]", "[Description]", "[Where used]", "Section 4.2"],
            ], [1.6, 2.6, 1.4, 0.9]),
        ]),
    ],
)

# --------------------------------------------------------------------------
# Independent Contractor Agreement (Luwah hires a subcontractor)
# --------------------------------------------------------------------------

COMPANY_SIGNER = luwah_signer("Company")

CONTRACTOR_SIGNER = [
    ("Signature", ""),
    ("Name", "[Contractor name]"),
    ("Title", "[Title, or Owner]"),
    ("Email", "[Contractor email]"),
    ("Date", ""),
    ("Notice address", "[Contractor notice address]"),
]

ICA = Doc(
    stem="Contractor_Agreement_LuwahTech_Template_v1",
    title="Independent Contractor Agreement",
    subtitle=f"Terms for subcontractors performing work for {COMPANY_SHORT} and its clients",
    summary=[
        ("Agreement No.", "[ICA-YYYY-NN]"),
        ("Contractor", "[Contractor legal name]"),
        ("Contractor entity type", "[Individual, LLC or corporation, and state]"),
        ("Contractor contact", "[Email and phone]"),
        ("Effective Date", "[Month DD, YYYY]"),
        ("Initial Work Order reference", "[WO-YYYY-NN, or None at signing]"),
        ("Compensation basis", "[Hourly at the Work Order rate, or fixed per Work Order]"),
    ],
    intro=[
        P(
            f"This Independent Contractor Agreement (the \"Agreement\") is between {COMPANY}, a "
            f"{STATE} limited liability company (\"Company\"), and the contractor named in the "
            "Agreement Summary above (\"Contractor\"). Company engages Contractor to perform "
            "services for Company and, at Company's direction, for Company's clients (each a "
            "\"Client\")."
        ),
        P(
            "Company owns the Client relationship and the work. Contractor works under Work Orders, "
            "controls how the work is done, and is responsible for its own taxes, tools and insurance."
        ),
    ],
    sections=[
        Section("Definitions", [
            CL("\"Work Order\"", "means a document in substantially the form of Exhibit A, signed by both Parties, that describes the Services, Deliverables, compensation and dates for an assignment."),
            CL("\"Services\"", "means the work Contractor performs under a Work Order."),
            CL("\"Work Product\"", "means everything Contractor creates, alone or with others, in performing the Services, including code, workflows, configurations, designs, content, documentation, drafts and data."),
            CL("\"Pre-Existing IP\"", "means tools, code, libraries and materials Contractor created or acquired before or outside the Services, listed in Exhibit B."),
            CL("\"Confidential Information\"", "means non-public information of Company or of any Client that Contractor receives or learns in connection with the Services, in any form. It includes Client identities, business data, credentials, systems, pricing, source code, Work Product and the terms of this Agreement."),
        ]),
        Section("Services", [
            C("Contractor performs the Services described in each Work Order. Each Work Order incorporates this Agreement."),
            C("Contractor decides how and when to perform the Services, subject to the deadlines, deliverable specifications and security requirements in the Work Order."),
            C("Contractor performs the Services personally, or through the named personnel listed in the Work Order. Contractor may not subcontract or delegate any part of the Services without Company's prior written consent."),
            C("This Agreement is non-exclusive. Company does not promise any minimum amount of work, and Contractor may work for others, subject to Sections 6, 7 and 8."),
            C("Where a Client's agreement with Company imposes terms on the work, such as data handling, security, confidentiality or ownership terms, Company will state them in the Work Order and Contractor will comply with them."),
        ]),
        Section("Compensation", [
            C("Company pays Contractor as stated in the Work Order, either hourly at the stated rate in 15-minute increments, or a fixed amount per Deliverable."),
            C("Contractor submits invoices monthly, or on completion of a milestone where the Work Order says so. Each invoice states the Work Order number and itemizes hours by date and task, or the Deliverables completed. Company pays approved invoices within 30 days of receipt (Net 30)."),
            C("Company may reject hours that were not authorized by the Work Order or are not itemized. The Parties resolve any rejected item in good faith. Company pays the undisputed part on time."),
            C("Company reimburses expenses only with prior written approval, at cost, with receipts."),
            C("Contractor is responsible for all taxes on amounts paid under this Agreement. Company issues Form 1099 where the law requires. Contractor provides a completed Form W-9, or Form W-8 if applicable, before the first payment."),
            C("Company does not pay for work performed without a signed Work Order."),
        ]),
        Section("Independent Contractor Status", [
            C("Contractor is an independent contractor and not an employee, partner, agent or joint venturer of Company. Contractor is not entitled to wages, overtime, benefits, workers' compensation or any other employee entitlement from Company."),
            C("Contractor supplies its own equipment, tools, software licenses and workspace. Company provides only the access to Company or Client systems that the Services require."),
            C("Contractor has no authority to bind Company or any Client, and will not represent itself as an employee or agent of either."),
            CLB("Notice required by Colorado law.", "CONTRACTOR IS NOT ENTITLED TO UNEMPLOYMENT INSURANCE BENEFITS UNLESS UNEMPLOYMENT COMPENSATION COVERAGE IS PROVIDED BY CONTRACTOR OR SOME OTHER ENTITY. CONTRACTOR IS OBLIGATED TO PAY FEDERAL AND STATE INCOME TAX ON ANY MONEYS PAID PURSUANT TO THIS AGREEMENT."),
        ]),
        Section("Work Product and Intellectual Property", [
            CL("Work made for hire.", "The Parties agree that all Work Product is specially ordered or commissioned by Company and is a \"work made for hire\" under the United States Copyright Act, 17 U.S.C. § 101, to the fullest extent the law allows, and that Company is its author and owner."),
            CL("Assignment.", "To the extent any Work Product does not qualify as a work made for hire, Contractor irrevocably assigns to Company, on creation, all right, title and interest in the Work Product worldwide, including all copyrights, and any patent, trademark and trade secret rights in it, together with the right to sue for past and future infringement."),
            CL("Moral rights.", "To the extent the law allows, Contractor waives, and agrees not to assert against Company, its Clients or their successors, any moral rights in the Work Product."),
            CL("Further assurances.", "Contractor will sign any document and take any lawful step Company reasonably requests, at Company's expense, to record or perfect Company's ownership, during and after the term."),
            CL("Pre-Existing IP.", "Contractor keeps ownership of Pre-Existing IP. Contractor will not incorporate Pre-Existing IP into Work Product unless it is listed in Exhibit B or approved in writing in the Work Order. For any Pre-Existing IP incorporated into Work Product, Contractor grants Company and its Clients a perpetual, worldwide, non-exclusive, royalty-free, transferable and sublicensable license to use, copy, modify and distribute it as part of the Work Product."),
            CL("Open source and third-party components.", "Contractor will not incorporate open source or third-party components into Work Product without Company's prior written approval, and will identify each component and its license."),
            CL("Flow-down to Clients.", "Company may assign or license Work Product to its Clients. Contractor's assignments, waivers and licenses in this Section extend to those Clients and their successors."),
            CL("Originality.", "Contractor warrants that the Work Product is Contractor's original work, except for approved Pre-Existing IP and approved third-party components, and does not infringe or misappropriate any third party's rights. Contractor will not use any confidential information or trade secret of a former client or employer in the Services."),
            CL("AI tools.", "Contractor may use AI-assisted tools only where the Work Order permits them, and only tools operated under terms that do not permit the vendor to train on submitted content. Contractor never submits Company or Client Confidential Information to a consumer tier of any AI service, never submits credentials to any AI tool, and submits protected health information, payment card data or other regulated data to an AI tool only where the Work Order expressly provides for it and states the written agreement that covers it. Contractor tells the participants before an AI tool records or transcribes a meeting. Contractor reviews all AI-assisted output and remains responsible for the Work Product."),
            CL("Company tools.", "Company's own tools and libraries, including the Luwah Accessibility Toolkit, the n8n workflow library, the web starter and the operations scripts, are Company property. Their source code, build configuration, design and documentation are Confidential Information even where a compiled or served copy is public. Contractor uses them only for the Work Order that requires them, does not copy them to any other system or client, keeps no copy after the Work Order ends, and does not build a similar product from their source, design or documentation during the term or for 12 months after it. This clause protects Company's trade secrets and Confidential Information. It is not a covenant not to compete and does not restrict Contractor from using general skills, knowledge and experience or publicly available information."),
        ]),
        Section("Confidentiality", [
            C("Contractor will use Confidential Information only to perform the Services, will protect it with at least reasonable care, and will not disclose it to anyone other than personnel named in the Work Order who are bound by written obligations at least as protective as this Section."),
            C("Confidential Information does not include information that is or becomes public through no fault of Contractor, that Contractor already knew without a duty of confidence, or that Contractor lawfully receives from a third party without a duty of confidence. Contractor may disclose Confidential Information when the law requires, after prompt notice to Company where lawful."),
            C("On Company's request, or when a Work Order or this Agreement ends, Contractor will return or delete all Confidential Information and Work Product within 5 business days and confirm in writing that it has done so, except copies the law requires Contractor to keep."),
            C("These obligations last for 3 years after this Agreement ends. For trade secrets and credentials, they last as long as the information remains a trade secret or the credential remains active."),
            CL("General knowledge.", GENERAL_KNOWLEDGE_CLAUSE.format(party="Contractor")),
            CL("Whistleblower immunity.", DTSA_NOTICE),
        ]),
        Section("Client Access, Credentials and Security", [
            C("Contractor uses only the access Company or a Client grants, only for the Work Order that requires it, and never for any other purpose."),
            C("Contractor stores every credential in an encrypted password manager, enables multi-factor authentication where the platform offers it, never shares credentials, never stores them in email, chat or source code, and never uses a personal account to access a Client system."),
            C("Contractor keeps Client data only in the systems the Work Order approves, never on an unencrypted device, and keeps every device used for the Services under full-disk encryption with current operating system and security updates."),
            C("Contractor makes no change to a Client production system outside the scope of the Work Order or outside the approval process Company sets for that Client."),
            C("Contractor reports any suspected security incident, including a lost device, a leaked credential or unauthorized access, to Company within 24 hours of learning of it, and cooperates fully with Company's and the Client's response."),
            C("When a Work Order ends, Contractor deletes all Client credentials and Client data it holds within 5 business days and confirms in writing that it has done so."),
            C("Company may check Contractor's compliance with this Section on reasonable notice."),
        ]),
        Section("Client Relationships", [
            C("Company owns the Client relationship. Contractor communicates with a Client only through the channels Company approves and only within the scope of the Work Order."),
            C("Contractor will not use Company's trade secrets or Confidential Information, including a Client's identity, needs, systems, pricing or contacts learned through the Services, to solicit or divert any Client. This clause protects Company's trade secrets and Confidential Information. It is not a covenant not to compete and does not stop Contractor from working for anyone using information Contractor did not obtain through the Services."),
            C("If a Client asks Contractor to work directly for it during a Work Order, Contractor will tell Company within 2 business days."),
            C("Contractor will not make any public statement that disparages Company or a Client."),
        ]),
        Section("Representations and Compliance", [
            C("Contractor has the right to enter this Agreement and to grant the rights in it, and no other agreement or obligation conflicts with it."),
            C("Contractor will comply with all laws that apply to the Services and to Contractor's business, and holds every license and registration its business requires."),
            C("Contractor is legally able to work in the country where it performs the Services."),
        ]),
        Section("Term and Termination", [
            C("This Agreement starts on the Effective Date and continues until either Party ends it."),
            C("Either Party may end this Agreement or a Work Order with 7 days' written notice. Company may end this Agreement or a Work Order immediately by written notice if Contractor breaches Section 5, 6, 7 or 8, or if a Client asks Company to remove Contractor from its work."),
            C("When this Agreement or a Work Order ends, Contractor delivers all Work Product completed or in progress, returns or deletes Confidential Information under Section 6.3, and deletes credentials under Section 7.6. Company pays for Services performed up to the end date, subject to Sections 3.3 and 11."),
            C("Sections 5, 6, 7, 8, 11, 12 and 14, and any other term that by its nature should survive, survive the end of this Agreement."),
        ]),
        Section("Warranty", [
            C("Contractor warrants that it will perform the Services in a professional manner consistent with industry standards, and that each Deliverable will conform to the Work Order."),
            C("For 30 days after Company accepts a Deliverable, Contractor will correct any non-conformance at no charge within a reasonable time. If Contractor cannot, Company may have the work corrected elsewhere and deduct the reasonable cost from amounts owed to Contractor."),
        ]),
        Section("Indemnity and Limitation of Liability", [
            C("Contractor will defend, indemnify and hold harmless Company and its Clients from any third-party claim, and the resulting losses, costs and reasonable attorneys' fees, that arises from Contractor's breach of Section 5, 6 or 7, from infringement or misappropriation by the Work Product, from Contractor's gross negligence or willful misconduct, or from any tax, wage or employment claim by Contractor or its personnel."),
            C("Company's total liability arising out of this Agreement is limited to the compensation paid to Contractor under it in the 12 months before the event that gave rise to the claim. Neither Party is liable for indirect, incidental, special, consequential or punitive damages, except for Contractor's breach of Section 5, 6 or 7 and Contractor's indemnity obligations."),
        ]),
        Section("Insurance", [
            C("Contractor is responsible for its own insurance. Where a Work Order requires professional liability or cyber liability insurance, Contractor will hold the stated coverage for the term of the Work Order and provide a certificate of insurance on request."),
        ]),
        Section("General", [
            CL("Governing law and disputes.", f"The law of the State of {STATE} governs this Agreement. The Parties will first negotiate any dispute in good faith for 30 days, then mediate in {COUNTY} or by video conference with a mediator they agree on, sharing the mediator's fees equally. If mediation fails, either Party may bring the dispute in the state or federal courts located in {COUNTY}, and each Party submits to their jurisdiction. Either Party may seek an injunction in any court to protect its Confidential Information or intellectual property without first negotiating or mediating. The prevailing Party may recover its reasonable attorneys' fees and costs."),
            NOTICES_CLAUSE,
            CL("Entire agreement.", "This Agreement, its Exhibits and each Work Order are the entire agreement between the Parties about their subject matter and replace all earlier proposals, discussions and agreements about it."),
            CL("Amendment and waiver.", "Changes must be in writing and signed by both Parties. A signature by email or through an e-signature platform counts. A Party's failure to enforce a term is not a waiver of it."),
            CL("Assignment.", "Contractor may not assign this Agreement or any Work Order. Company may assign this Agreement to a successor to substantially all of its business, on written notice."),
            SEVERABILITY_CLAUSE,
            HEADINGS_CLAUSE,
        ]),
    ],
    parties=("Company", "Contractor"),
    left_signer=COMPANY_SIGNER,
    right_signer=CONTRACTOR_SIGNER,
    doc_prefix="ICA",
    cover_tag="CONTRACTOR AGREEMENT   ·   SUBCONTRACT TERMS",
    cover_line="[Contractor name]   |   [Client and project, or Internal]",
    cover_parties=(f"{COMPANY} (\"Company\")", "[Contractor legal name] (\"Contractor\")"),
    exhibits=[
        Section("Exhibit A: Work Order", [
            N("One Work Order per assignment. Every Work Order incorporates the Independent Contractor Agreement. Work starts only after both Parties sign."),
            F([
                ("Work Order No.", "[WO-YYYY-NN]"),
                ("Contractor", "[Contractor legal name]"),
                ("Agreement reference", "[ICA-YYYY-NN dated Month DD, YYYY]"),
                ("Client and project", "[Client name and project, or Internal]"),
                ("Company SOW reference", "[SOW-YYYY-NN, or None]"),
                ("Start Date", "[Month DD, YYYY]"),
                ("End Date", "[Month DD, YYYY]"),
                ("Named personnel", "[Contractor personnel permitted to perform the Services, or Contractor only]"),
            ]),
            H("1. Services"),
            P("[What Contractor will do, on which systems, and what it will not do.]"),
            H("2. Deliverables and Acceptance Criteria"),
            T(["No.", "Deliverable", "Acceptance criteria", "Due date"], BLANK_ROWS, [0.5, 2.2, 2.8, 1.0]),
            H("3. Compensation"),
            T(["Basis", "Rate or amount (USD)", "Not-to-exceed (USD)", "Invoicing"], [
                ["[Hourly, 15-minute increments]", "[Rate per hour]", "[Cap for this Work Order]", "[Monthly, Net 30]"],
                ["[Fixed per Deliverable]", "[Amount per Deliverable]", "[Total]", "[On acceptance, Net 30]"],
            ], [1.9, 1.6, 1.5, 1.5]),
            H("4. Access to Be Granted"),
            T(["System", "Access level", "Granted by", "Revoked on"], [["", "", "", "End Date"], ["", "", "", "End Date"]], [2.0, 1.6, 1.4, 1.5]),
            H("5. Client Terms and Security Requirements"),
            N("State every term the Client's agreement with Company imposes on this work: data handling, regulated data, approval process for production changes, permitted tools, and insurance."),
            B(["[Client term]", "[Approval process for production changes]", "[Permitted AI tools, or None]", "[Required insurance, or None]"]),
            SIG("Company", "Contractor", COMPANY_SIGNER[:5], CONTRACTOR_SIGNER[:5]),
        ]),
        Section("Exhibit B: Contractor Pre-Existing IP", [
            N("List every item of Contractor Pre-Existing IP that may be incorporated into Work Product. Anything not listed here or approved in a Work Order may not be incorporated. If nothing applies, write NONE in the first row."),
            T(["Item", "Description", "Where it will be used", "License"], [["[Item, or NONE]", "", "", "Section 5.5"], ["", "", "", "Section 5.5"]], [1.6, 2.6, 1.4, 0.9]),
        ]),
    ],
)

# --------------------------------------------------------------------------
# Mutual Non-Disclosure Agreement (Vikunja #311)
# --------------------------------------------------------------------------

NDA = Doc(
    stem="NDA_Mutual_LuwahTech_Template_v1",
    title="Mutual Non-Disclosure Agreement",
    subtitle=f"How {COMPANY_SHORT} and a client protect each other's confidential information",
    summary=[
        ("Agreement No.", "[NDA-YYYY-NN]"),
        ("Client", "[Client legal name and entity type, or an individual doing business as Trade Name]"),
        ("Client contact", "[Name, title, email]"),
        ("Effective Date", "[Month DD, YYYY]"),
        ("Purpose", "[One line, for example: evaluating an n8n automation build for candidate intake]"),
        ("Term", "[2] years from the Effective Date. Section 6."),
        ("Related documents", "[MSA-YYYY-NN, proposal dated Month DD, YYYY, or None]"),
    ],
    intro=[
        P(
            f"This Mutual Non-Disclosure Agreement (the \"Agreement\") is between {COMPANY}, a "
            f"{STATE} limited liability company (\"Provider\"), and the client named in the "
            "Agreement Summary above (\"Client\"). Provider and Client are each a \"Party\" and "
            "together the \"Parties\". The Agreement takes effect on the Effective Date."
        ),
        P(
            "The Parties expect to share confidential information with each other while they "
            "discuss, evaluate and carry out the Purpose stated in the Agreement Summary. This "
            "Agreement is mutual. The same obligations apply to both Parties, whichever one is "
            "disclosing and whichever one is receiving. It also covers Confidential Information "
            "the Parties made available to each other in connection with the Purpose before the "
            "Effective Date."
        ),
    ],
    sections=[
        Section("Definitions", [
            CL("\"Confidential Information\"", "means information that one Party (the \"Disclosing Party\") or its Representatives makes available to the other Party (the \"Receiving Party\"), in any form, that is marked or identified as confidential, or that a reasonable person would understand to be confidential from its nature or from the way it was disclosed. It includes business plans, pricing and rates, client, customer and candidate lists, financial information, technical information, source code, system configurations, credentials and access to systems, workflow and automation designs, data, and the existence and terms of any proposal or agreement between the Parties."),
            CL("Oral and visual disclosure.", "Information disclosed orally, visually, in a meeting or on a shared screen is Confidential Information on the same test. Neither Party needs to confirm it in writing afterwards for the protection to apply."),
            CL("\"Representatives\"", "means a Party's members, managers, officers, employees, subcontractors and professional advisers who need the Confidential Information for the Purpose."),
            CL("\"Trade secret\"", "has the meaning given in the Colorado Uniform Trade Secrets Act and the federal Defend Trade Secrets Act."),
            CL("\"Purpose\"", "means the purpose stated in the Agreement Summary."),
        ]),
        Section("What Is Not Confidential Information", [
            C("Confidential Information does not include information that the Receiving Party can demonstrate:"),
            B([
                "is or becomes publicly available, other than through a breach of this Agreement by the Receiving Party or its Representatives;",
                "was lawfully known to the Receiving Party before the Disclosing Party disclosed it, free of any duty of confidentiality;",
                "is lawfully received from a third party who is not under a duty of confidentiality to the Disclosing Party; or",
                "is independently developed by the Receiving Party without using the Disclosing Party's Confidential Information.",
            ]),
            C("A combination of details is not public just because each detail is separately public, unless the combination itself is public."),
        ]),
        Section("Obligations of the Receiving Party", [
            C("The Receiving Party uses Confidential Information only for the Purpose."),
            C("The Receiving Party protects Confidential Information with at least the care it uses for its own confidential information, and never less than reasonable care."),
            C("The Receiving Party shares Confidential Information only with Representatives who need it for the Purpose and who are bound by written confidentiality obligations at least as protective as this Agreement, or by a professional duty of confidentiality. The Receiving Party is responsible for what its Representatives do with it."),
            C("The Receiving Party shares Confidential Information with anyone else only with the Disclosing Party's prior written consent."),
            C("The Receiving Party does not copy Confidential Information except as the Purpose requires. Every copy carries the same protection as the original."),
            C("The Receiving Party does not reverse engineer, decompile or disassemble any software, prototype or device the Disclosing Party provides."),
            C("The Receiving Party tells the Disclosing Party promptly on learning of any unauthorized use or disclosure, and cooperates to contain it."),
            CL("Credentials and system access.", "Where Confidential Information includes login credentials, API keys or access to a system, the Receiving Party uses that access only for the Purpose, keeps the credentials in an encrypted password manager, does not share them outside its Representatives, and returns or stops using the access under Section 7. Access granted for the Purpose is not a license to the system."),
            CL("AI tools.", AI_CONFIDENTIALITY_CLAUSE),
        ]),
        Section("Personal Data and Regulated Information", [
            C("Confidential Information may include personal information about third parties, for example the Client's customers, patients, candidates or employees. Each Party handles that information in line with the laws that apply to it, and shares only what the Purpose requires."),
            CL("Protected health information.", "Protected health information under HIPAA is not shared under this Agreement. If the Purpose requires it, the Parties sign a business associate agreement first. Where that agreement and this one conflict about protected health information, the business associate agreement governs. If protected health information reaches the Receiving Party despite this Section, the Receiving Party treats it as Confidential Information, does not use it, tells the Disclosing Party promptly, and returns or destroys it on request."),
            C("Neither Party shares payment card numbers, government identification numbers or account passwords belonging to third parties under this Agreement unless the Purpose requires it and the Parties have agreed in writing how they will be protected."),
        ]),
        Section("Naming Each Other", [
            C("Either Party may identify the other by name as a client or vendor, and describe the general nature of the work, in its portfolio and marketing only with the other Party's prior written consent, which may be given by email and withdrawn at any time by written notice. A name or description disclosed under that consent is not Confidential Information. The terms of any proposal or agreement between the Parties, and everything else the consent does not cover, stay Confidential Information."),
        ]),
        Section("Term and Survival", [
            C("This Agreement starts on the Effective Date and runs for the Term stated in the Agreement Summary. Either Party may end it earlier with [30] days' written notice."),
            C("Ending this Agreement does not release either Party from its obligations for Confidential Information disclosed before the end. Those obligations continue for [3] years after the Confidential Information was disclosed, or [3] years after this Agreement ends, whichever is later. Ending this Agreement early does not shorten them."),
            C("For Confidential Information that is a trade secret, and for credentials, the obligations continue for as long as the information remains a trade secret or the credential remains active."),
            C("If the Parties later sign a Master Services Agreement, Work Made for Hire Agreement or other agreement with its own confidentiality terms, the more protective confidentiality terms govern the work under that agreement. This Agreement continues to govern everything else the Parties share."),
            C("Every Section of this Agreement survives its end for as long as any obligation under it continues."),
        ]),
        Section("Return, Deletion and Access", [
            C("On the Disclosing Party's written request, or when this Agreement ends, the Receiving Party, within [15] days, returns or destroys the Disclosing Party's Confidential Information and every copy, at the Disclosing Party's choice."),
            C("Within the same period the Receiving Party stops using any access it was granted, returns any credentials it holds, and confirms that it no longer has access."),
            C("The Receiving Party confirms in writing that it has done so, if asked."),
            C("The Receiving Party may keep copies held in routine backup systems that are not readily accessible, and copies it must keep to meet a legal, regulatory or professional record-keeping obligation. Retained copies stay confidential under this Agreement for as long as they are held, and are not used for any other purpose."),
        ]),
        Section("Disclosure Required by Law", [
            C("The Receiving Party may disclose Confidential Information where a law, regulation, subpoena, court order or government request requires it. Where the law allows, the Receiving Party gives the Disclosing Party prompt written notice before disclosing, cooperates with the Disclosing Party's reasonable efforts to limit or protect the disclosure, at the Disclosing Party's expense, and discloses only the part that is legally required. Information disclosed this way remains Confidential Information for every other purpose."),
        ]),
        Section("Ownership, No License, No Warranty", [
            C("Confidential Information remains the property of the Disclosing Party. This Agreement grants no license or other right in it, except the limited right to use it for the Purpose."),
            C("Confidential Information is provided as is. The Disclosing Party gives no warranty about its accuracy, completeness or fitness for any purpose, except that each Party confirms it has the right to disclose what it discloses."),
            C("Nothing in this Agreement obliges either Party to disclose any information, to enter into any further agreement, or to continue discussions. Either Party may end discussions at any time. Any further work between the Parties is covered by a separate written agreement."),
        ]),
        Section("General Knowledge and Skills", [
            C("Each Party may continue to use the general knowledge, skills and experience it gains through the Purpose, including methods and techniques of general application, provided it does not use or disclose the other Party's Confidential Information in doing so. This Section grants no license under any patent, copyright, trade secret or other intellectual property right of the other Party."),
        ]),
        Section("Remedies", [
            C("A breach of this Agreement may cause harm that money cannot fully repair. The Disclosing Party may seek an injunction or other equitable relief to prevent or stop a breach, without posting a bond where the court allows, in addition to any other remedy available to it."),
        ]),
        Section("Whistleblower Notice", [
            C(DTSA_NOTICE),
        ]),
        Section("General", [
            CL("Relationship.", "The Parties are independent. This Agreement creates no partnership, joint venture, agency or employment relationship, and neither Party may bind the other."),
            CL("Assignment.", "Neither Party may assign this Agreement without the other Party's written consent, except to a successor to substantially all of its business, on written notice to the other Party."),
            NOTICES_CLAUSE,
            CL("Entire agreement.", "This Agreement is the entire agreement between the Parties about its subject matter and replaces any earlier discussion of confidentiality between them."),
            CL("Amendment and waiver.", "Changes must be in writing and signed by both Parties. A signature by email or through an e-signature platform counts. A Party's failure to enforce a term, or delay in enforcing it, is not a waiver of it."),
            CL("Governing law and disputes.", f"The law of the State of {STATE} governs this Agreement, without regard to its conflict of law rules. The Parties will first negotiate any dispute in good faith for 30 days, then mediate in {COUNTY} or by video conference with a mediator they agree on, sharing the mediator's fees equally. If mediation fails, either Party may bring the dispute in the state or federal courts located in {COUNTY}, and each Party submits to their jurisdiction. Either Party may seek an injunction in any court with jurisdiction to protect its Confidential Information without first negotiating or mediating. Where a Master Services Agreement between the Parties is in force, its dispute process applies to a confidentiality claim that also arises under it. The prevailing Party may recover its reasonable attorneys' fees and costs."),
            SEVERABILITY_CLAUSE,
            HEADINGS_CLAUSE,
        ]),
    ],
    parties=("Provider", "Client"),
    left_signer=PROVIDER_SIGNER,
    right_signer=CLIENT_SIGNER,
    exhibits=[],
    doc_prefix="NDA",
    cover_tag="CLIENT AGREEMENT   ·   MUTUAL NDA",
)

DOCS = [NDA, MSA, WFH, ICA]

# --------------------------------------------------------------------------
# DOCX rendering
# --------------------------------------------------------------------------

PLACEHOLDER = re.compile(r"(\[[^\]]+\])")


def set_font(run, name: str, size: float | None = None, bold: bool | None = None,
             color: RGBColor | None = None, italic: bool | None = None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_text(paragraph, text: str, size: float = 10.5, bold: bool = False,
             color: RGBColor = INK, italic: bool = False):
    """Write text, coloring [bracketed placeholders] copper so they are easy to find."""
    for part in PLACEHOLDER.split(text):
        if not part:
            continue
        run = paragraph.add_run(part)
        if PLACEHOLDER.fullmatch(part):
            set_font(run, BODY_FONT, size, True, COPPER, italic)
        else:
            set_font(run, BODY_FONT, size, bold, color, italic)


def shade(cell, hex_fill: str):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def bottom_border(paragraph, hex_color: str, size: int = 8):
    ppr = paragraph._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), hex_color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def add_field(run, instr: str):
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._element.append(el)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
            run._element.append(el)
        else:
            el = OxmlElement("w:t")
            el.text = text
            run._element.append(el)


def mark_header_row(row):
    """Flag a table's first row as a header row. Screen readers associate the
    column names with each cell, and Word repeats the row across page breaks."""
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:tblHeader")) is None:
        trpr.append(OxmlElement("w:tblHeader"))


def describe_picture(run, description: str):
    """Alt text for an inline picture. python-docx exposes no API for it."""
    for docpr in run._element.iter(qn("wp:docPr")):
        docpr.set("descr", description)
        docpr.set("title", description)


def set_cell_widths(table, widths_in: list[float]):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            row.cells[idx].width = Inches(width)


def cell_text(cell, text: str, bold: bool = False, size: float = 9.5, color: RGBColor = INK):
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.space_before = Pt(2)
    add_text(para, text, size=size, bold=bold, color=color)


def new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    section.top_margin = Inches(1.1)
    section.bottom_margin = Inches(0.9)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)

    for name, size, color in (("Heading 1", 12.5, RGBColor(0xFF, 0xFF, 0xFF)), ("Heading 2", 11.5, BLUE), ("Heading 3", 10.5, INK)):
        style = doc.styles[name]
        style.font.name = HEAD_FONT
        style.font.size = Pt(size)
        style.font.bold = name != "Heading 1"
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(6 if name == "Heading 1" else 4)
        style.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            # Section band in the house Letter of Agreement style: white on copper.
            ppr = style.element.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "B87333")
            ppr.append(shd)
            style.paragraph_format.left_indent = Inches(0.12)
            pbdr = OxmlElement("w:pBdr")
            for side in ("top", "bottom"):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "12")
                el.set(qn("w:space"), "3")
                el.set(qn("w:color"), "B87333")
                pbdr.append(el)
            ppr.append(pbdr)
        srpr = style.element.get_or_add_rPr()
        srfonts = srpr.find(qn("w:rFonts"))
        if srfonts is None:
            srfonts = OxmlElement("w:rFonts")
            srpr.append(srfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            srfonts.set(qn(attr), HEAD_FONT)
        # Drop the theme font reference so Word does not override the named font.
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if srfonts.get(qn(attr)) is not None:
                del srfonts.attrib[qn(attr)]

    # Header: logo on the left, template version on the right, blue rule below.
    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(4)
    logo_run = hp.add_run()
    logo_run.add_picture(str(LOGO), width=Inches(2.3))
    describe_picture(logo_run, "Luwah Technologies LLC logo")
    tab_run = hp.add_run("\t\t")
    ver = hp.add_run(f"Template {VERSION}")
    set_font(ver, HEAD_FONT, 9, False, MUTED)
    set_font(tab_run, HEAD_FONT, 9)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    bottom_border(hp, "4A90A4", 6)

    # Footer: contact line, then page X of Y.
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run(f"{COMPANY}  |  {CITY_LINE}  |  {EMAIL}  |  {PHONE}  |  {WEB}")
    set_font(r, BODY_FONT, 7.5, False, MUTED)
    fp2 = footer.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp2.paragraph_format.space_after = Pt(0)
    r1 = fp2.add_run(f"{TAGLINE}  |  Page ")
    set_font(r1, BODY_FONT, 7.5, False, MUTED)
    r2 = fp2.add_run()
    set_font(r2, BODY_FONT, 7.5, False, MUTED)
    add_field(r2, "PAGE")
    r3 = fp2.add_run(" of ")
    set_font(r3, BODY_FONT, 7.5, False, MUTED)
    r4 = fp2.add_run()
    set_font(r4, BODY_FONT, 7.5, False, MUTED)
    add_field(r4, "NUMPAGES")
    return doc


def render_form(doc, fields: list[tuple[str, str]]):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in fields:
        row = table.add_row()
        cell_text(row.cells[0], label, bold=True)
        shade(row.cells[0], SHADE)
        cell_text(row.cells[1], value)
    set_cell_widths(table, [2.0, 4.5])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_table(doc, header: list[str], rows: list[list[str]], widths: list[float] | None):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        cell_text(cell, text, bold=True, size=9, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "4A90A4")
    mark_header_row(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        for idx, text in enumerate(row_values):
            cell_text(row.cells[idx], text, size=9)
    if widths:
        set_cell_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_signature(doc, left: str, right: str, left_rows, right_rows):
    heading = doc.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(10)
    add_text(heading, "Signatures", size=11, bold=True, color=BLUE)
    intro = doc.add_paragraph()
    intro.paragraph_format.keep_with_next = True
    add_text(intro, "Each Party signs below to agree to these terms. An electronic signature has the same effect as an original.")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for idx, title in enumerate((left, right)):
        cell = table.rows[0].cells[idx]
        cell_text(cell, title, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "4A90A4")
    mark_header_row(table.rows[0])
    for (llabel, lvalue), (rlabel, rvalue) in zip(left_rows, right_rows):
        row = table.add_row()
        for cell, label, value in ((row.cells[0], llabel, lvalue), (row.cells[1], rlabel, rvalue)):
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(6 if label == "Signature" else 2)
            para.paragraph_format.space_after = Pt(2)
            add_text(para, f"{label}: ", size=9, bold=True, color=MUTED)
            if label == "Signature":
                add_text(para, "\n\n______________________________", size=9)
            else:
                add_text(para, value, size=9.5)
    set_cell_widths(table, [3.25, 3.25])
    # Keep the whole signature block on one page: no row may split, and every
    # row but the last asks Word to keep it with the next.
    for row_idx, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        if row_idx < len(table.rows) - 1:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True
    doc.add_paragraph()


def render_blocks(doc, blocks, section_no: int | None):
    clause = 0
    for block in blocks:
        kind = block[0]
        if kind in ("c", "cl", "clb"):
            clause += 1
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.45)
            para.paragraph_format.first_line_indent = Inches(-0.45)
            para.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
            number = f"{section_no}.{clause}\t" if section_no else f"{clause}.\t"
            add_text(para, number, bold=True, color=BLUE)
            if kind == "cl":
                add_text(para, block[1] + " ", bold=True)
                add_text(para, block[2])
            elif kind == "clb":
                add_text(para, block[1] + " ", bold=True)
                add_text(para, block[2], bold=True)
            else:
                add_text(para, block[1])
        elif kind == "b":
            for item in block[1]:
                para = doc.add_paragraph(style="List Bullet")
                para.paragraph_format.left_indent = Inches(0.8)
                para.paragraph_format.first_line_indent = Inches(-0.25)
                para.paragraph_format.space_after = Pt(3)
                add_text(para, item)
        elif kind == "p":
            para = doc.add_paragraph()
            add_text(para, block[1])
        elif kind == "n":
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.15)
            bottom_border(para, "D9E2E6", 4)
            add_text(para, block[1], size=9.5, color=MUTED, italic=True)
        elif kind == "h":
            # Level 2 under the level 1 exhibit title, so no heading level is skipped.
            doc.add_heading(block[1], level=2)
        elif kind == "t":
            render_table(doc, block[1], block[2], block[3])
        elif kind == "f":
            render_form(doc, block[1])
        elif kind == "sig":
            render_signature(doc, block[1], block[2], block[3], block[4])
        else:
            raise ValueError(f"unknown block {kind}")


def render_cover(doc, spec: Doc):
    """Dark cover page in the house Letter of Agreement style: tag line, thin
    wordmark, document title, client line, Between and Document blocks, tagline.
    One shaded single-cell table fills the page. Word paints cell shading edge to
    edge, which a paragraph background cannot do."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    shade(cell, "1C1C1E")
    tcpr = cell._element.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", 900), ("start", 720), ("bottom", 720), ("end", 720)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tcpr.append(margins)
    row = table.rows[0]
    trpr = row._tr.get_or_add_trPr()
    # "atLeast", never "exact": an exact row clips whatever wraps past it, and
    # the two longer covers lost the closing tagline that way.
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(int(8.4 * 1440)))
    height.set(qn("w:hRule"), "atLeast")
    trpr.append(height)

    white = RGBColor(0xFF, 0xFF, 0xFF)
    soft = RGBColor(0xCC, 0xCC, 0xCC)

    def line(text: str, font: str, size: float, color: RGBColor, bold: bool = False,
             before: float = 0, after: float = 0, spacing: int | None = None,
             align=WD_ALIGN_PARAGRAPH.LEFT):
        para = cell.add_paragraph()
        para.alignment = align
        para.paragraph_format.space_before = Pt(before)
        para.paragraph_format.space_after = Pt(after)
        run = para.add_run(text)
        set_font(run, font, size, bold, color)
        if spacing is not None:
            rpr = run._element.get_or_add_rPr()
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(spacing))
            rpr.append(sp)
        return para

    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    tag = line(spec.cover_tag, HEAD_FONT, 9, BLUE, spacing=50, after=6)
    bottom_border(tag, "B87333", 18)
    line("LUWAH", HEAD_FONT, 30, white, before=36, after=0, spacing=40)
    line("TECHNOLOGIES", HEAD_FONT, 30, COPPER, after=0, spacing=40)
    line(spec.title, HEAD_FONT, 24, white, before=84, after=0)
    sub = line(spec.subtitle, HEAD_FONT, 13, COPPER, after=4)
    bottom_border(sub, "555555", 4)
    line(spec.cover_line, BODY_FONT, 10.5, soft, before=6, after=0)
    line("BETWEEN", HEAD_FONT, 8.5, BLUE, before=84, after=2, spacing=40)
    for party in spec.cover_parties:
        line(party, BODY_FONT, 10, soft, after=0)
    line("DOCUMENT", HEAD_FONT, 8.5, BLUE, before=12, after=2, spacing=40)
    line(f"[{spec.doc_prefix}-YYYY-NN]   |   [Month DD, YYYY]   |   Template {VERSION}", BODY_FONT, 10, soft, after=0)
    line(TAGLINE, BODY_FONT, 10, COPPER, before=40, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)


def build_docx(spec: Doc, out_dir: Path) -> Path:
    doc = new_document()
    render_cover(doc, spec)

    title = doc.add_paragraph()
    # The title is the first paragraph after the cover table. Breaking before it
    # keeps the cover to one page. A break paragraph after the table spills over.
    title.paragraph_format.page_break_before = True
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run(spec.title)
    set_font(run, HEAD_FONT, 24, True, COPPER)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(10)
    run = sub.add_run(spec.subtitle)
    set_font(run, HEAD_FONT, 11, False, BLUE)

    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(3)
    label.paragraph_format.keep_with_next = True
    add_text(label, "Agreement Summary", size=10, bold=True, color=BLUE)
    render_form(doc, spec.summary)

    note = doc.add_paragraph()
    note.paragraph_format.left_indent = Inches(0.15)
    bottom_border(note, "D9E2E6", 4)
    add_text(note, TEMPLATE_NOTE, size=9.5, color=MUTED, italic=True)

    render_blocks(doc, spec.intro, None)

    for idx, section in enumerate(spec.sections, start=1):
        doc.add_heading(f"{idx}. {section.title}", level=1)
        render_blocks(doc, section.blocks, idx)

    render_signature(doc, spec.parties[0], spec.parties[1], spec.left_signer, spec.right_signer)

    for exhibit in spec.exhibits:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(exhibit.title, level=1)
        render_blocks(doc, exhibit.blocks, None)

    core = doc.core_properties
    core.language = "en-US"
    core.title = f"{spec.title} - {COMPANY} Template {VERSION}"
    core.author = COMPANY
    core.subject = spec.subtitle
    core.keywords = "Luwah Technologies, template, agreement"
    core.comments = f"Template {VERSION} built {VERSION_DATE} by scripts/build-agreements.py"

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{spec.stem}.docx"
    doc.save(path)
    return path


WORD_CONTAINER = Path.home() / "Library/Containers/com.microsoft.Word/Data/Documents"


def export_pdf(docx_path: Path) -> Path:
    """Export through Microsoft Word on macOS.

    Word is sandboxed. From AppleScript it can only write inside its own container,
    so the PDF is saved there first and then moved next to the DOCX. After a
    "save as" to PDF the document reference changes, so close the active document.
    """
    import subprocess

    pdf_path = docx_path.with_suffix(".pdf")
    # Open from inside the container too. Opening from any other path makes
    # Word raise a "Grant File Access" dialog that only a human can dismiss.
    # A per-run subdirectory keeps two runs from unlinking each other's files.
    staging = WORD_CONTAINER / f"build-{os.getpid()}"
    staging.mkdir(parents=True, exist_ok=True)
    staged_docx = staging / docx_path.name
    staged_pdf = staging / pdf_path.name
    for stale in (staged_docx, staged_pdf):
        if stale.exists():
            stale.unlink()
    shutil.copy2(docx_path, staged_docx)
    # Word's "close" only accepts "active document", so check its name before
    # every close. The owner may have an unsaved Word window in front, and
    # "saving no" would discard it silently if the wrong document were active.
    script = f'''
tell application "Microsoft Word"
  open file name (POSIX file "{staged_docx}" as string)
  if name of active document is not "{docx_path.name}" then error "expected {docx_path.name} to be active after open"
  save as active document file name "{staged_pdf}" file format format PDF
  if name of active document is not "{docx_path.name}" then error "expected {docx_path.name} to be active after save"
  close active document saving no
end tell
'''
    try:
        result = subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=180)
        if result.returncode != 0 or not staged_pdf.exists():
            raise RuntimeError(f"PDF export failed for {docx_path.name}: {result.stderr.strip()}")
        shutil.move(str(staged_pdf), str(pdf_path))
    finally:
        shutil.rmtree(staging, ignore_errors=True)
    return pdf_path


class WordLock:
    """One export run at a time on this Mac. Word serves "active document" to
    whichever script asks, so two runs at once export each other's files. A
    count of open documents is a check, not a lock; this is the lock."""

    def __init__(self):
        WORD_CONTAINER.mkdir(parents=True, exist_ok=True)
        self.path = WORD_CONTAINER / "build-agreements.lock"
        self.handle = None

    def __enter__(self):
        import fcntl

        self.handle = open(self.path, "w")
        try:
            fcntl.flock(self.handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            self.handle.close()
            raise RuntimeError("another build-agreements export holds the Word lock; wait for it to finish") from exc
        self.handle.write(str(os.getpid()))
        self.handle.flush()
        return self

    def __exit__(self, *exc):
        import fcntl

        fcntl.flock(self.handle, fcntl.LOCK_UN)
        self.handle.close()
        return False


def word_open_documents() -> int:
    """How many documents Word has open. Zero when Word is not running."""
    import subprocess

    script = '''
if application "Microsoft Word" is running then
  tell application "Microsoft Word" to return count of documents
end if
return 0
'''
    result = subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=60)
    if result.returncode != 0:
        raise RuntimeError(f"could not query Word: {result.stderr.strip()}")
    return int(result.stdout.strip() or "0")


def quit_word() -> None:
    """Quit Word so the next run opens fresh files. Only called when nothing is open,
    so "saving no" cannot discard anyone's work. Word re-serves a same-named document
    it already has open, which silently exports a stale PDF."""
    import subprocess

    subprocess.run(["osascript", "-e", 'tell application "Microsoft Word" to quit saving no'],
                   capture_output=True, text=True, timeout=60)


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, default=OUT_DEFAULT, help="output directory (default public/agreements)")
    parser.add_argument("--copy-to", type=Path, default=None, help="also copy the outputs into this directory")
    parser.add_argument("--no-pdf", action="store_true", help="skip the PDF export through Word")
    parser.add_argument("--only", choices=[d.stem for d in DOCS], default=None)
    args = parser.parse_args(argv)

    if not LOGO.exists():
        print(f"logo missing: {LOGO}", file=sys.stderr)
        return 2

    specs = [spec for spec in DOCS if not args.only or spec.stem == args.only]
    outputs: list[Path] = []

    if args.no_pdf:
        for spec in specs:
            docx_path = build_docx(spec, args.out)
            outputs.append(docx_path)
            print(f"wrote {docx_path.relative_to(ROOT) if docx_path.is_relative_to(ROOT) else docx_path}")
    else:
        # Hold the lock for the whole run, and render nothing until it is held,
        # so a refused run leaves no new DOCX beside an old PDF.
        with WordLock():
            open_docs = word_open_documents()
            if open_docs:
                print(f"Microsoft Word has {open_docs} document(s) open. Close them, then rerun. "
                      "An open document with a template's name would export stale.", file=sys.stderr)
                return 3
            for spec in specs:
                docx_path = build_docx(spec, args.out)
                outputs.append(docx_path)
                print(f"wrote {docx_path.relative_to(ROOT) if docx_path.is_relative_to(ROOT) else docx_path}")
                pdf_path = export_pdf(docx_path)
                outputs.append(pdf_path)
                print(f"wrote {pdf_path.relative_to(ROOT) if pdf_path.is_relative_to(ROOT) else pdf_path}")
            if word_open_documents() == 0:
                quit_word()

    # The page reads this so the site and the documents state one version.
    version_file = ROOT / "src" / "app" / "agreements" / "version.json"
    version_file.write_text(
        json.dumps({"version": VERSION, "date": VERSION_DATE, "label": VERSION_LABEL}, indent=2) + "\n"
    )
    print(f"wrote {version_file.relative_to(ROOT)}")

    if args.copy_to:
        args.copy_to.mkdir(parents=True, exist_ok=True)
        for path in outputs:
            shutil.copy2(path, args.copy_to / path.name)
            print(f"copied {path.name} -> {args.copy_to}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
